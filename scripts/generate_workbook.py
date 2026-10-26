#!/usr/bin/env python3
"""
hemanthx446 Company Intelligence - Excel Workbook Generator
Generates a multi-tab, professionally styled company_intelligence.xlsx workbook
using openpyxl, adhering to management-grade design standards and strict data contracts.
"""

import sys
import os
import json
import argparse
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Design Palette - Professional Executive Navy/Slate
COLOR_PRIMARY_NAVY = "1B365D"   # Primary headers
COLOR_SECONDARY_BLUE = "2B547E" # Subheaders / section titles
COLOR_ACCENT_BLUE = "4A90E2"    # Highlights
COLOR_LIGHT_ZEBRA = "F4F7FA"    # Alternating data rows
COLOR_WHITE = "FFFFFF"
COLOR_CARD_BG = "EBF2FA"        # KPI summary cards
COLOR_BORDER = "D1D5DB"
COLOR_TEXT_MAIN = "1F2937"

# Standard Fonts
FONT_TITLE = Font(name="Segoe UI", size=16, bold=True, color=COLOR_WHITE)
FONT_SECTION = Font(name="Segoe UI", size=12, bold=True, color=COLOR_PRIMARY_NAVY)
FONT_SUBHEADER = Font(name="Segoe UI", size=11, bold=True, color=COLOR_WHITE)
FONT_HEADER = Font(name="Segoe UI", size=10, bold=True, color=COLOR_WHITE)
FONT_BOLD = Font(name="Segoe UI", size=10, bold=True, color=COLOR_TEXT_MAIN)
FONT_REGULAR = Font(name="Segoe UI", size=10, color=COLOR_TEXT_MAIN)
FONT_SMALL = Font(name="Segoe UI", size=9, italic=True, color="6B7280")
FONT_LINK = Font(name="Segoe UI", size=10, color="0000FF", underline="single")
FONT_CARD_NUM = Font(name="Segoe UI", size=16, bold=True, color=COLOR_PRIMARY_NAVY)
FONT_CARD_LBL = Font(name="Segoe UI", size=9, bold=True, color="4B5563")

# Fills
FILL_PRIMARY = PatternFill(start_color=COLOR_PRIMARY_NAVY, end_color=COLOR_PRIMARY_NAVY, fill_type="solid")
FILL_SECONDARY = PatternFill(start_color=COLOR_SECONDARY_BLUE, end_color=COLOR_SECONDARY_BLUE, fill_type="solid")
FILL_ZEBRA = PatternFill(start_color=COLOR_LIGHT_ZEBRA, end_color=COLOR_LIGHT_ZEBRA, fill_type="solid")
FILL_WHITE = PatternFill(start_color=COLOR_WHITE, end_color=COLOR_WHITE, fill_type="solid")
FILL_CARD = PatternFill(start_color=COLOR_CARD_BG, end_color=COLOR_CARD_BG, fill_type="solid")

# Borders
BORDER_THIN = Border(
    left=Side(style='thin', color=COLOR_BORDER),
    right=Side(style='thin', color=COLOR_BORDER),
    top=Side(style='thin', color=COLOR_BORDER),
    bottom=Side(style='thin', color=COLOR_BORDER)
)
BORDER_TOP_THICK = Border(
    top=Side(style='medium', color=COLOR_PRIMARY_NAVY),
    bottom=Side(style='thin', color=COLOR_BORDER)
)

# Alignments
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)

def apply_table_styling(ws, header_row, start_col, end_col, num_rows):
    """Applies frozen panes, header styling, alternating fills, and borders."""
    ws.freeze_panes = ws[f"A{header_row + 1}"]
    
    # Header row
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = FONT_HEADER
        cell.fill = FILL_PRIMARY
        cell.alignment = ALIGN_HEADER
        cell.border = BORDER_THIN
    ws.row_dimensions[header_row].height = 28

    # Data rows
    for r in range(header_row + 1, header_row + 1 + num_rows):
        is_zebra = ((r - header_row) % 2 == 0)
        row_fill = FILL_ZEBRA if is_zebra else FILL_WHITE
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=col)
            if cell.fill.fill_type is None:
                cell.fill = row_fill
            cell.border = BORDER_THIN
            if cell.alignment.horizontal is None:
                cell.alignment = ALIGN_LEFT
        ws.row_dimensions[r].height = 22

def auto_fit_columns(ws, min_width=12, max_width=50):
    """Calculates appropriate column widths based on content length."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.value is not None:
                # Truncate string representation for line breaks
                lines = str(cell.value).split("\n")
                line_max = max(len(l) for l in lines)
                if line_max > max_len:
                    max_len = line_max
        calc_width = max(min_width, min(max_len + 3, max_width))
        ws.column_dimensions[col_letter].width = calc_width

def build_executive_summary_sheet(ws, data):
    """Constructs the Executive Summary dashboard sheet."""
    ws.title = "Executive Summary"
    ws.views.sheetView[0].showGridLines = True
    
    meta = data.get("metadata", {})
    exec_data = data.get("executive_summary", {})
    
    # 1. Header Banner
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = f"COMPANY INTELLIGENCE EXECUTIVE DOSSIER: {meta.get('target_company_name', 'TARGET ENTERPRISE').upper()}"
    title_cell.font = FONT_TITLE
    title_cell.fill = FILL_PRIMARY
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 40

    # Sub-banner metadata
    ws.merge_cells("A2:H2")
    sub_cell = ws["A2"]
    sub_cell.value = (
        f"Project: {meta.get('project_name', 'hemanthx446 Company Intelligence')}  |  "
        f"Investigator: {meta.get('owner', 'Hemanth')}  |  "
        f"Date: {meta.get('research_date', datetime.today().strftime('%Y-%m-%d'))}  |  "
        f"Confidence: {meta.get('overall_confidence', 'Medium')}  |  "
        f"Website: {meta.get('target_website', 'N/A')}"
    )
    sub_cell.font = FONT_SUBHEADER
    sub_cell.fill = FILL_SECONDARY
    sub_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 24

    # 2. KPI Summary Cards (Row 4-5)
    counts = {
        "Offerings": len(data.get("offerings", [])),
        "Confirmed Customers": len(data.get("customers", [])),
        "Official Partners": len(data.get("partners", [])),
        "Key Contacts": len(data.get("contacts", [])),
        "Signals Detected": len(data.get("business_signals", [])),
        "Opportunities": len(data.get("opportunities", [])),
    }
    
    # Top score calculation
    opps = data.get("opportunities", [])
    top_score = max([o.get("total_score", 0) for o in opps]) if opps else 0
    counts["Top Opp Score"] = f"{top_score}/100"

    card_positions = [("A", "Offerings"), ("B", "Confirmed Customers"), ("C", "Official Partners"), 
                      ("D", "Key Contacts"), ("E", "Signals Detected"), ("F", "Opportunities"), ("G", "Top Opp Score")]

    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 28

    for col_letter, label in card_positions:
        val = counts[label]
        # Label cell
        lbl_cell = ws[f"{col_letter}4"]
        lbl_cell.value = label.upper()
        lbl_cell.font = FONT_CARD_LBL
        lbl_cell.fill = FILL_CARD
        lbl_cell.alignment = ALIGN_CENTER
        lbl_cell.border = BORDER_THIN

        # Number cell
        num_cell = ws[f"{col_letter}5"]
        num_cell.value = val
        num_cell.font = FONT_CARD_NUM
        num_cell.fill = FILL_CARD
        num_cell.alignment = ALIGN_CENTER
        num_cell.border = BORDER_THIN

    # 3. Executive Overview & Strategic Narrative
    ws.cell(row=7, column=1, value="1. Strategic Synthesis & Commercial Brief").font = FONT_SECTION
    
    overview_text = exec_data.get(
        "overview", 
        "Comprehensive open-source corporate intelligence research synthesis. "
        "Evaluates technological footprints, confirmed commercial ecosystems, key organizational stakeholders, "
        "and prioritized transformation opportunities."
    )
    ws.merge_cells("A8:G9")
    narrative_cell = ws["A8"]
    narrative_cell.value = overview_text
    narrative_cell.font = FONT_REGULAR
    narrative_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    narrative_cell.border = BORDER_THIN
    ws.row_dimensions[8].height = 28
    ws.row_dimensions[9].height = 28

    # 4. Top Priority Commercial Opportunities Table
    ws.cell(row=11, column=1, value="2. Prioritized Commercial Opportunities").font = FONT_SECTION
    
    headers_opp = ["Opp ID", "Opportunity Title & Core Scope", "Fit (30)", "Need (25)", "Access (20)", "Comm (15)", "Feas (10)", "Total", "Confidence"]
    for c_idx, h in enumerate(headers_opp, start=1):
        c = ws.cell(row=12, column=c_idx, value=h)
        c.font = FONT_HEADER
        c.fill = FILL_PRIMARY
        c.alignment = ALIGN_HEADER
        c.border = BORDER_THIN
    ws.row_dimensions[12].height = 25

    sorted_opps = sorted(opps, key=lambda x: x.get("total_score", 0), reverse=True)[:5]
    for r_idx, opp in enumerate(sorted_opps, start=13):
        is_zebra = (r_idx % 2 == 0)
        row_fill = FILL_ZEBRA if is_zebra else FILL_WHITE
        row_vals = [
            opp.get("id", f"OPP-{r_idx-12:03d}"),
            opp.get("title", "Commercial Opportunity"),
            opp.get("strategic_fit_score", 0),
            opp.get("need_evidence_score", 0),
            opp.get("contact_access_score", 0),
            opp.get("commercial_relevance_score", 0),
            opp.get("feasibility_timing_score", 0),
            opp.get("total_score", 0),
            opp.get("confidence", "Medium")
        ]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = FONT_REGULAR
            cell.fill = row_fill
            cell.border = BORDER_THIN
            if c_idx in [1, 8, 9]:
                cell.alignment = ALIGN_CENTER
                if c_idx == 8:
                    cell.font = FONT_BOLD
            elif c_idx >= 3 and c_idx <= 7:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT
        ws.row_dimensions[r_idx].height = 22

    # 5. Key Decision Makers Alignment
    next_start = 13 + len(sorted_opps) + 1
    ws.cell(row=next_start, column=1, value="3. Primary Stakeholder & Decision-Maker Alignment").font = FONT_SECTION
    
    headers_contacts = ["Contact ID", "Executive Name", "Current Title / Role", "Functional Area", "Decision Role", "Profile Link"]
    for c_idx, h in enumerate(headers_contacts, start=1):
        c = ws.cell(row=next_start+1, column=c_idx, value=h)
        c.font = FONT_HEADER
        c.fill = FILL_PRIMARY
        c.alignment = ALIGN_HEADER
        c.border = BORDER_THIN
    ws.row_dimensions[next_start+1].height = 25

    contacts = data.get("contacts", [])[:4]
    for r_offset, contact in enumerate(contacts, start=next_start+2):
        is_zebra = (r_offset % 2 == 0)
        row_fill = FILL_ZEBRA if is_zebra else FILL_WHITE
        
        ws.cell(row=r_offset, column=1, value=contact.get("id", f"CONT-{r_offset:03d}")).alignment = ALIGN_CENTER
        ws.cell(row=r_offset, column=2, value=contact.get("name", "N/A")).alignment = ALIGN_LEFT
        ws.cell(row=r_offset, column=3, value=contact.get("title", "N/A")).alignment = ALIGN_LEFT
        ws.cell(row=r_offset, column=4, value=contact.get("functional_area", "N/A")).alignment = ALIGN_LEFT
        ws.cell(row=r_offset, column=5, value=contact.get("decision_making_role", "Evaluator")).alignment = ALIGN_LEFT
        
        link_cell = ws.cell(row=r_offset, column=6, value="View Profile")
        url = contact.get("profile_url", "")
        if url and url.startswith("http"):
            link_cell.hyperlink = url
            link_cell.font = FONT_LINK
        else:
            link_cell.value = "Not disclosed"
            link_cell.font = FONT_REGULAR
        link_cell.alignment = ALIGN_CENTER

        for c_idx in range(1, 7):
            c = ws.cell(row=r_offset, column=c_idx)
            c.fill = row_fill
            c.border = BORDER_THIN
            if c_idx != 6:
                c.font = FONT_REGULAR
        ws.row_dimensions[r_offset].height = 22

    auto_fit_columns(ws)

def build_tabular_sheet(ws, title, headers, rows_data, url_col_indices=None):
    """Generic high-fidelity tabular sheet builder with auto-formatting and filters."""
    ws.title = title
    ws.views.sheetView[0].showGridLines = True
    
    # 1. Sheet Header Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=title.upper())
    title_cell.font = FONT_TITLE
    title_cell.fill = FILL_PRIMARY
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 35

    # 2. Table Column Headers
    header_row = 3
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_SECONDARY
        cell.alignment = ALIGN_HEADER
        cell.border = BORDER_THIN
    ws.row_dimensions[header_row].height = 28

    # 3. Data Rows
    num_rows = len(rows_data)
    for r_idx, row in enumerate(rows_data, start=header_row + 1):
        is_zebra = ((r_idx - header_row) % 2 == 0)
        row_fill = FILL_ZEBRA if is_zebra else FILL_WHITE
        
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = FONT_REGULAR
            cell.fill = row_fill
            cell.border = BORDER_THIN
            
            # Format lists or joins
            if isinstance(val, list):
                val_str = ", ".join(str(item) for item in val) if val else "None"
                cell.value = val_str
            else:
                cell.value = val if val is not None else "Unknown"

            # Check if this column should be a hyperlink
            if url_col_indices and c_idx in url_col_indices and isinstance(val, str) and val.startswith("http"):
                cell.hyperlink = val
                cell.font = FONT_LINK
                cell.alignment = ALIGN_LEFT
            elif isinstance(val, (int, float)):
                cell.alignment = ALIGN_CENTER
            elif str(val).startswith(("PROF-", "OFF-", "CUST-", "PART-", "CONT-", "COMP-", "SIG-", "OPP-", "OUT-", "SRC-", "GAP-")):
                cell.alignment = ALIGN_CENTER
            elif val in ["Verified fact", "Corroborated finding", "Reasonable inference", "Unverified claim", "Historical information", "Unknown information"]:
                cell.alignment = ALIGN_CENTER
            elif val in ["High", "Medium", "Low"]:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT
                
        ws.row_dimensions[r_idx].height = 24

    # Apply Auto-Filter & Freeze Panes
    last_col_letter = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{header_row + max(num_rows, 1)}"
    ws.freeze_panes = ws[f"A{header_row + 1}"]

    auto_fit_columns(ws)

def generate_workbook(data, output_filepath):
    """Builds the complete 12-worksheet company intelligence Excel workbook."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # 1. Executive Summary
    ws_exec = wb.create_sheet("Executive Summary")
    build_executive_summary_sheet(ws_exec, data)

    # 2. Company Profile
    ws_profile = wb.create_sheet("Company Profile")
    headers_prof = ["Record_ID", "Attribute", "Value", "Category", "Evidence_Status", "Source_References", "Notes"]
    rows_prof = []
    for item in data.get("company_profile", []):
        rows_prof.append([
            item.get("id", ""),
            item.get("attribute", ""),
            item.get("value", ""),
            item.get("category", ""),
            item.get("evidence_status", "Verified fact"),
            item.get("sources", []),
            item.get("notes", "")
        ])
    build_tabular_sheet(ws_profile, "Company Profile", headers_prof, rows_prof)

    # 3. Offerings
    ws_offerings = wb.create_sheet("Offerings")
    headers_off = ["Offering_ID", "Product_Solution_Name", "Category", "Technology_Stack", "Value_Proposition", "Target_Market", "Evidence_Status", "Source_References"]
    rows_off = []
    for item in data.get("offerings", []):
        rows_off.append([
            item.get("id", ""),
            item.get("name", ""),
            item.get("category", ""),
            item.get("tech_stack", ""),
            item.get("value_proposition", ""),
            item.get("target_market", ""),
            item.get("evidence_status", "Verified fact"),
            item.get("sources", [])
        ])
    build_tabular_sheet(ws_offerings, "Offerings", headers_off, rows_off)

    # 4. Customer Accounts
    ws_customers = wb.create_sheet("Customer Accounts")
    headers_cust = ["Customer_ID", "Customer_Name", "Industry_Vertical", "Project_Scope", "Technologies_Used", "Documented_Outcomes", "Evidence_Status", "Confidence", "Source_References"]
    rows_cust = []
    for item in data.get("customers", []):
        rows_cust.append([
            item.get("id", ""),
            item.get("customer_name", ""),
            item.get("industry", ""),
            item.get("project_scope", ""),
            item.get("technologies_used", ""),
            item.get("documented_outcomes", ""),
            item.get("evidence_status", "Verified fact"),
            item.get("confidence", "High"),
            item.get("sources", [])
        ])
    build_tabular_sheet(ws_customers, "Customer Accounts", headers_cust, rows_cust)

    # 5. Partners
    ws_partners = wb.create_sheet("Partners")
    headers_part = ["Partner_ID", "Partner_Name", "Partner_Type", "Strategic_Focus", "Verification_Method", "Evidence_Status", "Source_References"]
    rows_part = []
    for item in data.get("partners", []):
        rows_part.append([
            item.get("id", ""),
            item.get("partner_name", ""),
            item.get("partner_type", ""),
            item.get("strategic_focus", ""),
            item.get("verification_method", ""),
            item.get("evidence_status", "Verified fact"),
            item.get("sources", [])
        ])
    build_tabular_sheet(ws_partners, "Partners", headers_part, rows_part)

    # 6. Key Contacts
    ws_contacts = wb.create_sheet("Key Contacts")
    headers_cont = ["Contact_ID", "Full_Name", "Current_Title", "Functional_Area", "Decision_Role", "Public_Profile_URL", "Source_References"]
    rows_cont = []
    for item in data.get("contacts", []):
        rows_cont.append([
            item.get("id", ""),
            item.get("name", ""),
            item.get("title", ""),
            item.get("functional_area", ""),
            item.get("decision_making_role", "Evaluator"),
            item.get("profile_url", ""),
            item.get("sources", [])
        ])
    build_tabular_sheet(ws_contacts, "Key Contacts", headers_cont, rows_cont, url_col_indices=[6])

    # 7. Market and Competitors
    ws_competitors = wb.create_sheet("Market and Competitors")
    headers_comp = ["Competitor_ID", "Competitor_Name", "Market_Overlap", "Target_Strengths", "Competitor_Strengths", "Vulnerabilities", "Source_References"]
    rows_comp = []
    for item in data.get("competitors", []):
        rows_comp.append([
            item.get("id", ""),
            item.get("competitor_name", ""),
            item.get("market_overlap", ""),
            item.get("target_strengths", "Specialized focus and agile delivery"),
            item.get("strengths", ""),
            item.get("vulnerabilities", ""),
            item.get("sources", [])
        ])
    build_tabular_sheet(ws_competitors, "Market and Competitors", headers_comp, rows_comp)

    # 8. Business Signals
    ws_signals = wb.create_sheet("Business Signals")
    headers_sig = ["Signal_ID", "Signal_Category", "Event_Description", "Date_Detected", "Commercial_Significance", "Source_References"]
    rows_sig = []
    for item in data.get("business_signals", []):
        rows_sig.append([
            item.get("id", ""),
            item.get("signal_category", ""),
            item.get("event_description", ""),
            item.get("date_detected", ""),
            item.get("commercial_impact", ""),
            item.get("sources", [])
        ])
    build_tabular_sheet(ws_signals, "Business Signals", headers_sig, rows_sig)

    # 9. Opportunities
    ws_opps = wb.create_sheet("Opportunities")
    headers_opp = [
        "Opportunity_ID", "Opportunity_Title", "Fit_Score_30", "Need_Score_25", "Access_Score_20", 
        "Commercial_Score_15", "Feasibility_Score_10", "Total_Score_100", "Confidence", "Analytical_Rationale", 
        "Related_Contacts", "Related_Signals", "Source_References"
    ]
    rows_opp = []
    for item in data.get("opportunities", []):
        rows_opp.append([
            item.get("id", ""),
            item.get("title", ""),
            item.get("strategic_fit_score", 0),
            item.get("need_evidence_score", 0),
            item.get("contact_access_score", 0),
            item.get("commercial_relevance_score", 0),
            item.get("feasibility_timing_score", 0),
            item.get("total_score", 0),
            item.get("confidence", "Medium"),
            item.get("rationale", ""),
            item.get("related_contacts", []),
            item.get("related_signals", []),
            item.get("sources", [])
        ])
    build_tabular_sheet(ws_opps, "Opportunities", headers_opp, rows_opp)

    # 10. Outreach Plan
    ws_outreach = wb.create_sheet("Outreach Plan")
    headers_out = ["Plan_ID", "Target_Contact_ID", "Opportunity_Ref", "Strategic_Value_Angle", "Suggested_Asset", "Channel", "Sequence_Stage", "Status"]
    rows_out = []
    for item in data.get("outreach_plan", []):
        rows_out.append([
            item.get("id", ""),
            item.get("target_contact_id", ""),
            item.get("opportunity_id", ""),
            item.get("value_angle", ""),
            item.get("suggested_asset", ""),
            item.get("channel", "Professional Network"),
            item.get("sequence_stage", "Initial Discovery"),
            item.get("status", "Planned")
        ])
    build_tabular_sheet(ws_outreach, "Outreach Plan", headers_out, rows_out)

    # 11. Source Register
    ws_sources = wb.create_sheet("Source Register")
    headers_src = ["Source_ID", "Document_Title_or_Publisher", "Direct_Source_URL", "Publication_Date", "Access_Date", "Evidence_Category", "Confidence_Level", "Summary_Notes"]
    rows_src = []
    for item in data.get("sources", []):
        rows_src.append([
            item.get("id", ""),
            item.get("title", ""),
            item.get("url", ""),
            item.get("publication_date", ""),
            item.get("access_date", ""),
            item.get("evidence_category", "Verified fact"),
            item.get("confidence", "High"),
            item.get("notes", "")
        ])
    build_tabular_sheet(ws_sources, "Source Register", headers_src, rows_src, url_col_indices=[3])

    # 12. Research Gaps
    ws_gaps = wb.create_sheet("Research Gaps")
    headers_gap = ["Gap_ID", "Information_Dimension", "Unresolved_Question", "Limitation_Reason", "Next_Action_Plan"]
    rows_gap = []
    for item in data.get("research_gaps", []):
        rows_gap.append([
            item.get("id", ""),
            item.get("dimension", ""),
            item.get("description", ""),
            item.get("cause", ""),
            item.get("next_action", "")
        ])
    build_tabular_sheet(ws_gaps, "Research Gaps", headers_gap, rows_gap)

    # Ensure output parent directory exists
    os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
    wb.save(output_filepath)
    print(f"[SUCCESS] Successfully generated workbook with 12 worksheets at: {output_filepath}")
    return wb

def main():
    parser = argparse.ArgumentParser(description="hemanthx446 Company Intelligence Excel Generator")
    parser.add_argument("--input", "-i", required=True, help="Path to input structured JSON research file (e.g. data.json)")
    parser.add_argument("--output", "-o", default="outputs/company_intelligence.xlsx", help="Path for generated company_intelligence.xlsx")
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"[ERROR] Input file not found: {args.input}", file=sys.stderr)
        sys.exit(1)

    with open(args.input, "r", encoding="utf-8") as f:
        try:
            data = json.load(f)
        except json.JSONDecodeError as e:
            print(f"[ERROR] Failed to parse JSON: {e}", file=sys.stderr)
            sys.exit(1)

    generate_workbook(data, args.output)

if __name__ == "__main__":
    main()
