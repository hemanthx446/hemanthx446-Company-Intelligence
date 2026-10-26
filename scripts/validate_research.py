#!/usr/bin/env python3
"""
hemanthx446 Company Intelligence - Research Validation Script
Validates data integrity, schema compliance, source citation integrity,
scoring consistency, and generated Excel workbook completeness.
"""

import sys
import os
import json
import argparse
from datetime import datetime
import openpyxl

ALLOWED_EVIDENCE_STATUSES = {
    "Verified fact",
    "Corroborated finding",
    "Reasonable inference",
    "Unverified claim",
    "Historical information",
    "Unknown information"
}

ALLOWED_CONFIDENCE_LEVELS = {"High", "Medium", "Low"}

REQUIRED_TOP_LEVEL_KEYS = [
    "metadata",
    "executive_summary",
    "company_profile",
    "offerings",
    "customers",
    "partners",
    "contacts",
    "competitors",
    "business_signals",
    "opportunities",
    "outreach_plan",
    "sources",
    "research_gaps"
]

EXPECTED_WORKSHEETS = [
    "Executive Summary",
    "Company Profile",
    "Offerings",
    "Customer Accounts",
    "Partners",
    "Key Contacts",
    "Market and Competitors",
    "Business Signals",
    "Opportunities",
    "Outreach Plan",
    "Source Register",
    "Research Gaps"
]

class ResearchValidator:
    def __init__(self, data_path, workbook_path=None):
        self.data_path = data_path
        self.workbook_path = workbook_path
        self.errors = []
        self.warnings = []
        self.data = None

    def log_error(self, msg):
        self.errors.append(f"[ERROR] {msg}")

    def log_warning(self, msg):
        self.warnings.append(f"[WARNING] {msg}")

    def validate_all(self):
        print(f"--- Starting Validation on {self.data_path} ---")
        if not self.load_data():
            return False

        self.validate_metadata()
        self.validate_sources()
        self.validate_entities()
        self.validate_opportunities()
        self.validate_outreach_plan()
        
        if self.workbook_path:
            self.validate_workbook()

        self.print_summary()
        return len(self.errors) == 0

    def load_data(self):
        if not os.path.exists(self.data_path):
            self.log_error(f"Data file does not exist: {self.data_path}")
            return False
        try:
            with open(self.data_path, "r", encoding="utf-8") as f:
                self.data = json.load(f)
        except Exception as e:
            self.log_error(f"Failed to load or parse JSON: {e}")
            return False

        # Check required keys
        for key in REQUIRED_TOP_LEVEL_KEYS:
            if key not in self.data:
                self.log_error(f"Missing required top-level key: '{key}'")
        return True

    def validate_metadata(self):
        meta = self.data.get("metadata", {})
        required_fields = ["project_name", "owner", "target_company_name", "target_website", "target_country", "target_industry", "research_date"]
        for f in required_fields:
            val = meta.get(f)
            if not val or val == "N/A":
                self.log_error(f"Metadata field '{f}' is missing or empty.")

        if meta.get("project_name") != "hemanthx446 Company Intelligence":
            self.log_warning(f"Unexpected project name: '{meta.get('project_name')}'. Expected 'hemanthx446 Company Intelligence'.")

        if meta.get("owner") != "Hemanth":
            self.log_warning(f"Unexpected owner: '{meta.get('owner')}'. Expected 'Hemanth'.")

    def validate_sources(self):
        sources = self.data.get("sources", [])
        if not sources:
            self.log_error("Source register is empty. Every intelligence dossier must contain source citations.")
            return

        seen_ids = set()
        seen_urls = set()
        for idx, src in enumerate(sources):
            src_id = src.get("id")
            if not src_id:
                self.log_error(f"Source at index {idx} has no 'id'.")
            elif src_id in seen_ids:
                self.log_error(f"Duplicate Source ID detected: '{src_id}'")
            else:
                seen_ids.add(src_id)

            url = src.get("url", "")
            if not url or not url.startswith("http"):
                self.log_error(f"Source '{src_id}' has invalid URL: '{url}'")
            elif url in seen_urls:
                self.log_warning(f"Duplicate Source URL registered: '{url}' under ID '{src_id}'")
            else:
                seen_urls.add(url)

            cat = src.get("evidence_category")
            if cat not in ALLOWED_EVIDENCE_STATUSES:
                self.log_error(f"Source '{src_id}' has invalid evidence_category: '{cat}'")

            conf = src.get("confidence")
            if conf not in ALLOWED_CONFIDENCE_LEVELS:
                self.log_error(f"Source '{src_id}' has invalid confidence: '{conf}'")

    def validate_entities(self):
        valid_source_ids = {s.get("id") for s in self.data.get("sources", []) if s.get("id")}

        entity_configs = [
            ("company_profile", "PROF-", ["id", "attribute", "value"]),
            ("offerings", "OFF-", ["id", "name", "category"]),
            ("customers", "CUST-", ["id", "customer_name", "industry"]),
            ("partners", "PART-", ["id", "partner_name", "partner_type"]),
            ("contacts", "CONT-", ["id", "name", "title"]),
            ("competitors", "COMP-", ["id", "competitor_name", "market_overlap"]),
            ("business_signals", "SIG-", ["id", "signal_category", "event_description"]),
            ("research_gaps", "GAP-", ["id", "dimension", "description"])
        ]

        for entity_key, prefix, req_fields in entity_configs:
            items = self.data.get(entity_key, [])
            seen_ids = set()
            for idx, item in enumerate(items):
                item_id = item.get("id")
                if not item_id:
                    self.log_error(f"Item in '{entity_key}' at index {idx} has no 'id'.")
                    continue
                if not item_id.startswith(prefix):
                    self.log_warning(f"Item ID '{item_id}' in '{entity_key}' does not follow expected prefix '{prefix}'.")
                if item_id in seen_ids:
                    self.log_error(f"Duplicate ID '{item_id}' in '{entity_key}'.")
                seen_ids.add(item_id)

                for rf in req_fields:
                    if not item.get(rf):
                        self.log_error(f"Item '{item_id}' in '{entity_key}' missing required field: '{rf}'")

                # Validate evidence status if present
                ev_status = item.get("evidence_status")
                if ev_status and ev_status not in ALLOWED_EVIDENCE_STATUSES:
                    self.log_error(f"Item '{item_id}' in '{entity_key}' has invalid evidence_status: '{ev_status}'")

                # Validate source citations
                sources = item.get("sources", [])
                if not isinstance(sources, list):
                    self.log_error(f"Item '{item_id}' in '{entity_key}' 'sources' must be a list of source IDs.")
                else:
                    for src_ref in sources:
                        if src_ref not in valid_source_ids:
                            self.log_error(f"Item '{item_id}' in '{entity_key}' references non-existent source: '{src_ref}'")

    def validate_opportunities(self):
        opps = self.data.get("opportunities", [])
        if not opps:
            self.log_warning("No opportunities defined in research dataset.")
            return

        valid_contacts = {c.get("id") for c in self.data.get("contacts", []) if c.get("id")}
        valid_signals = {s.get("id") for s in self.data.get("business_signals", []) if s.get("id")}
        valid_sources = {s.get("id") for s in self.data.get("sources", []) if s.get("id")}

        seen_ids = set()
        for opp in opps:
            opp_id = opp.get("id")
            if not opp_id:
                self.log_error("Opportunity missing 'id'.")
                continue
            if opp_id in seen_ids:
                self.log_error(f"Duplicate Opportunity ID: '{opp_id}'")
            seen_ids.add(opp_id)

            # Check individual component weights
            fit = opp.get("strategic_fit_score", 0)
            need = opp.get("need_evidence_score", 0)
            access = opp.get("contact_access_score", 0)
            comm = opp.get("commercial_relevance_score", 0)
            feas = opp.get("feasibility_timing_score", 0)
            total = opp.get("total_score", 0)

            if not (0 <= fit <= 30):
                self.log_error(f"Opportunity '{opp_id}' strategic_fit_score ({fit}) out of range [0, 30].")
            if not (0 <= need <= 25):
                self.log_error(f"Opportunity '{opp_id}' need_evidence_score ({need}) out of range [0, 25].")
            if not (0 <= access <= 20):
                self.log_error(f"Opportunity '{opp_id}' contact_access_score ({access}) out of range [0, 20].")
            if not (0 <= comm <= 15):
                self.log_error(f"Opportunity '{opp_id}' commercial_relevance_score ({comm}) out of range [0, 15].")
            if not (0 <= feas <= 10):
                self.log_error(f"Opportunity '{opp_id}' feasibility_timing_score ({feas}) out of range [0, 10].")

            expected_total = fit + need + access + comm + feas
            if total != expected_total:
                self.log_error(f"Opportunity '{opp_id}' total_score ({total}) does not match sum of components ({expected_total}).")

            conf = opp.get("confidence")
            if conf not in ALLOWED_CONFIDENCE_LEVELS:
                self.log_error(f"Opportunity '{opp_id}' has invalid confidence: '{conf}'")

            # Check linked contacts & signals
            for c_ref in opp.get("related_contacts", []):
                if c_ref not in valid_contacts:
                    self.log_warning(f"Opportunity '{opp_id}' references unknown contact: '{c_ref}'")
            for s_ref in opp.get("related_signals", []):
                if s_ref not in valid_signals:
                    self.log_warning(f"Opportunity '{opp_id}' references unknown signal: '{s_ref}'")
            for src_ref in opp.get("sources", []):
                if src_ref not in valid_sources:
                    self.log_error(f"Opportunity '{opp_id}' references non-existent source: '{src_ref}'")

    def validate_outreach_plan(self):
        outreach = self.data.get("outreach_plan", [])
        valid_contacts = {c.get("id") for c in self.data.get("contacts", []) if c.get("id")}
        valid_opps = {o.get("id") for o in self.data.get("opportunities", []) if o.get("id")}

        seen_ids = set()
        for item in outreach:
            o_id = item.get("id")
            if not o_id:
                self.log_error("Outreach item missing 'id'.")
                continue
            if o_id in seen_ids:
                self.log_error(f"Duplicate Outreach Plan ID: '{o_id}'")
            seen_ids.add(o_id)

            c_id = item.get("target_contact_id")
            if c_id and c_id not in valid_contacts:
                self.log_warning(f"Outreach plan '{o_id}' references unknown contact: '{c_id}'")

            opp_id = item.get("opportunity_id")
            if opp_id and opp_id not in valid_opps:
                self.log_warning(f"Outreach plan '{o_id}' references unknown opportunity: '{opp_id}'")

    def validate_workbook(self):
        print(f"--- Validating Excel Workbook: {self.workbook_path} ---")
        if not os.path.exists(self.workbook_path):
            self.log_error(f"Workbook file does not exist: {self.workbook_path}")
            return

        try:
            wb = openpyxl.load_workbook(self.workbook_path, data_only=True)
        except Exception as e:
            self.log_error(f"Failed to load Excel workbook with openpyxl: {e}")
            return

        sheet_names = wb.sheetnames
        # Check all 12 sheets
        for expected in EXPECTED_WORKSHEETS:
            if expected not in sheet_names:
                self.log_error(f"Workbook missing required sheet: '{expected}'")

        # Check total number of sheets
        if len(sheet_names) != 12:
            self.log_warning(f"Workbook has {len(sheet_names)} sheets, expected exactly 12.")

        # Check Executive Summary content
        if "Executive Summary" in sheet_names:
            ws = wb["Executive Summary"]
            title_val = str(ws["A1"].value or "")
            target_name = self.data.get("metadata", {}).get("target_company_name", "").upper()
            if target_name and target_name not in title_val:
                self.log_warning(f"Executive Summary title banner '{title_val}' does not contain target company name '{target_name}'.")

    def print_summary(self):
        print("\n================ VALIDATION SUMMARY ================")
        print(f"Total Errors  : {len(self.errors)}")
        print(f"Total Warnings: {len(self.warnings)}")
        
        if self.warnings:
            print("\nWarnings:")
            for w in self.warnings:
                print(f"  {w}")

        if self.errors:
            print("\nErrors:")
            for e in self.errors:
                print(f"  {e}")
            print("\nRESULT: FAILED [X]")
        else:
            print("\nRESULT: PASSED [V] - Data and workbook integrity verified successfully.")
        print("====================================================\n")

def main():
    parser = argparse.ArgumentParser(description="Validate Company Intelligence Research Data")
    parser.add_argument("--input", "-i", required=True, help="Path to structured research JSON file (e.g. data.json)")
    parser.add_argument("--workbook", "-w", help="Optional path to generated Excel workbook for verification")
    args = parser.parse_args()

    validator = ResearchValidator(args.input, args.workbook)
    success = validator.validate_all()
    sys.exit(0 if success else 1)

if __name__ == "__main__":
    main()
