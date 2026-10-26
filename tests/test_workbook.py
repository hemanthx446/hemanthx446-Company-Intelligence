"""
Unit and integration tests for hemanthx446 Company Intelligence research system.
Verifies workbook generation, schema validation, 12 worksheets, hyperlinks,
scoring calculations, missing data resilience, and referential integrity.
"""

import os
import sys
import json
import pytest
import openpyxl

# Add scripts directory to path
SCRIPTS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "scripts"))
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

from generate_workbook import generate_workbook
from validate_research import ResearchValidator, EXPECTED_WORKSHEETS

TEST_DATA_PATH = os.path.join(os.path.dirname(__file__), "test_data", "synthetic_company_data.json")
TEMP_WORKBOOK_PATH = os.path.join(os.path.dirname(__file__), "test_output", "test_company_intelligence.xlsx")

@pytest.fixture(scope="module")
def synthetic_data():
    """Loads the synthetic test dataset."""
    assert os.path.exists(TEST_DATA_PATH), f"Synthetic test data not found at {TEST_DATA_PATH}"
    with open(TEST_DATA_PATH, "r", encoding="utf-8") as f:
        return json.load(f)

@pytest.fixture(scope="module")
def generated_workbook(synthetic_data):
    """Generates a test workbook and returns its path."""
    os.makedirs(os.path.dirname(TEMP_WORKBOOK_PATH), exist_ok=True)
    generate_workbook(synthetic_data, TEMP_WORKBOOK_PATH)
    assert os.path.exists(TEMP_WORKBOOK_PATH)
    return TEMP_WORKBOOK_PATH

def test_generate_workbook_creates_all_12_sheets(generated_workbook):
    """Verifies that all 12 required worksheets exist in the workbook."""
    wb = openpyxl.load_workbook(generated_workbook, data_only=True)
    sheet_names = wb.sheetnames
    
    assert len(sheet_names) == 12, f"Expected 12 worksheets, found {len(sheet_names)}: {sheet_names}"
    for expected in EXPECTED_WORKSHEETS:
        assert expected in sheet_names, f"Missing required worksheet: '{expected}'"

def test_headers_present_on_all_tabular_sheets(generated_workbook):
    """Verifies that each tabular sheet has headers in row 3."""
    wb = openpyxl.load_workbook(generated_workbook, data_only=True)
    
    tabular_checks = {
        "Company Profile": ["Record_ID", "Attribute", "Value"],
        "Offerings": ["Offering_ID", "Product_Solution_Name", "Category"],
        "Customer Accounts": ["Customer_ID", "Customer_Name", "Industry_Vertical"],
        "Partners": ["Partner_ID", "Partner_Name", "Partner_Type"],
        "Key Contacts": ["Contact_ID", "Full_Name", "Current_Title"],
        "Market and Competitors": ["Competitor_ID", "Competitor_Name", "Market_Overlap"],
        "Business Signals": ["Signal_ID", "Signal_Category", "Event_Description"],
        "Opportunities": ["Opportunity_ID", "Opportunity_Title", "Total_Score_100"],
        "Outreach Plan": ["Plan_ID", "Target_Contact_ID", "Opportunity_Ref"],
        "Source Register": ["Source_ID", "Document_Title_or_Publisher", "Direct_Source_URL"],
        "Research Gaps": ["Gap_ID", "Information_Dimension", "Unresolved_Question"]
    }
    
    for sheet_name, expected_cols in tabular_checks.items():
        ws = wb[sheet_name]
        row_values = [ws.cell(row=3, column=col).value for col in range(1, ws.max_column + 1)]
        for exp in expected_cols:
            assert exp in row_values, f"Sheet '{sheet_name}' missing expected header '{exp}'. Found: {row_values}"

def test_hyperlinks_formatted(generated_workbook):
    """Verifies that URLs are formatted with clickable hyperlinks in openpyxl."""
    wb = openpyxl.load_workbook(generated_workbook)
    
    # Check Source Register sheet column 3 (Direct_Source_URL)
    ws_src = wb["Source Register"]
    # Row 4 is first data row
    url_cell = ws_src.cell(row=4, column=3)
    assert url_cell.hyperlink is not None, "Source Register URL cell has no hyperlink object"
    assert url_cell.hyperlink.target.startswith("http"), f"Invalid hyperlink target: {url_cell.hyperlink.target}"

    # Check Key Contacts sheet column 6 (Public_Profile_URL)
    ws_cont = wb["Key Contacts"]
    cont_url_cell = ws_cont.cell(row=4, column=6)
    assert cont_url_cell.hyperlink is not None, "Key Contacts profile URL cell has no hyperlink object"
    assert cont_url_cell.hyperlink.target.startswith("http")

def test_executive_summary_content(generated_workbook, synthetic_data):
    """Verifies that Executive Summary reflects actual metadata and records."""
    wb = openpyxl.load_workbook(generated_workbook, data_only=True)
    ws = wb["Executive Summary"]
    
    target_name = synthetic_data["metadata"]["target_company_name"].upper()
    title_banner = str(ws["A1"].value)
    assert target_name in title_banner, f"Expected '{target_name}' in executive summary banner '{title_banner}'"
    
    # Check KPI cards
    assert ws["A5"].value == len(synthetic_data["offerings"])
    assert ws["B5"].value == len(synthetic_data["customers"])
    assert ws["C5"].value == len(synthetic_data["partners"])
    assert ws["D5"].value == len(synthetic_data["contacts"])
    assert ws["F5"].value == len(synthetic_data["opportunities"])

def test_safe_handling_of_missing_data(tmp_path):
    """Ensures generator does not crash on empty or partially missing data structures."""
    minimal_data = {
        "metadata": {
            "project_name": "hemanthx446 Company Intelligence",
            "owner": "Hemanth",
            "target_company_name": "Minimal Test Corp",
            "target_website": "https://minimal.example.com",
            "target_country": "USA",
            "target_industry": "Software",
            "research_date": "2026-10-26"
        },
        "executive_summary": {},
        "company_profile": [],
        "offerings": [],
        "customers": [],
        "partners": [],
        "contacts": [],
        "competitors": [],
        "business_signals": [],
        "opportunities": [],
        "outreach_plan": [],
        "sources": [],
        "research_gaps": []
    }
    
    output_file = str(tmp_path / "minimal_workbook.xlsx")
    wb = generate_workbook(minimal_data, output_file)
    assert os.path.exists(output_file)
    assert len(wb.sheetnames) == 12

def test_validator_passes_valid_data(generated_workbook):
    """Validates that ResearchValidator passes clean test data without errors."""
    validator = ResearchValidator(TEST_DATA_PATH, generated_workbook)
    is_valid = validator.validate_all()
    assert is_valid, f"Validation failed with errors: {validator.errors}"
    assert len(validator.errors) == 0

def test_validator_catches_invalid_scoring_and_broken_sources(tmp_path):
    """Ensures validator detects broken scoring, invalid source links, and missing keys."""
    broken_data = {
        "metadata": {
            "project_name": "hemanthx446 Company Intelligence",
            "owner": "Hemanth",
            "target_company_name": "Broken Corp",
            "target_website": "https://broken.example.com",
            "target_country": "USA",
            "target_industry": "Tech",
            "research_date": "2026-10-26"
        },
        "executive_summary": {},
        "company_profile": [],
        "offerings": [],
        "customers": [],
        "partners": [],
        "contacts": [],
        "competitors": [],
        "business_signals": [],
        "opportunities": [
            {
                "id": "OPP-001",
                "title": "Bad Math Opportunity",
                "strategic_fit_score": 35,  # Exceeds max 30
                "need_evidence_score": 20,
                "contact_access_score": 15,
                "commercial_relevance_score": 10,
                "feasibility_timing_score": 5,
                "total_score": 999,  # Doesn't match sum
                "confidence": "InvalidConfidence",  # Not High/Med/Low
                "sources": ["SRC-NON-EXISTENT"]     # Broken citation
            }
        ],
        "outreach_plan": [],
        "sources": [
            {
                "id": "SRC-001",
                "title": "Valid Source",
                "url": "not-a-valid-url",
                "evidence_category": "InvalidCategory",
                "confidence": "High"
            }
        ],
        "research_gaps": []
    }
    
    broken_file = str(tmp_path / "broken_data.json")
    with open(broken_file, "w", encoding="utf-8") as f:
        json.dump(broken_data, f)
        
    validator = ResearchValidator(broken_file)
    is_valid = validator.validate_all()
    
    assert not is_valid, "Validator should have failed on corrupted data."
    assert any("out of range" in e for e in validator.errors)
    assert any("does not match sum" in e for e in validator.errors)
    assert any("non-existent source" in e for e in validator.errors)
    assert any("invalid evidence_category" in e for e in validator.errors)

def test_validator_catches_duplicate_ids(tmp_path):
    """Ensures validator catches duplicate entity and source IDs."""
    dup_data = {
        "metadata": {
            "project_name": "hemanthx446 Company Intelligence",
            "owner": "Hemanth",
            "target_company_name": "Dup Corp",
            "target_website": "https://dup.example.com",
            "target_country": "USA",
            "target_industry": "Tech",
            "research_date": "2026-10-26"
        },
        "executive_summary": {},
        "company_profile": [
            {"id": "PROF-001", "attribute": "Attr 1", "value": "Val 1"},
            {"id": "PROF-001", "attribute": "Attr 2", "value": "Val 2"}  # Duplicate PROF-001
        ],
        "offerings": [],
        "customers": [],
        "partners": [],
        "contacts": [],
        "competitors": [],
        "business_signals": [],
        "opportunities": [],
        "outreach_plan": [],
        "sources": [
            {"id": "SRC-001", "title": "S1", "url": "https://s1.com", "evidence_category": "Verified fact", "confidence": "High"},
            {"id": "SRC-001", "title": "S2", "url": "https://s2.com", "evidence_category": "Verified fact", "confidence": "High"}   # Duplicate SRC-001
        ],
        "research_gaps": []
    }
    
    dup_file = str(tmp_path / "dup_data.json")
    with open(dup_file, "w", encoding="utf-8") as f:
        json.dump(dup_data, f)
        
    validator = ResearchValidator(dup_file)
    is_valid = validator.validate_all()
    assert not is_valid
    assert any("Duplicate Source ID detected" in e for e in validator.errors)
    assert any("Duplicate ID 'PROF-001'" in e for e in validator.errors)
